import csv
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.analytics import selectors
from apps.analytics.filters import (
    AnalyticsFilters,
    parse_filters_payload,
    period_bounds,
)
from apps.analytics.models import AnalyticsExportJob, ExportType, FileFormat
from apps.boards.services import ColumnSettingsService
from apps.imports.models import ImportLog
from apps.jobs.models import BackgroundJob
from apps.notifications.models import NotificationJob
from apps.tasks.models import Task

ExportHandler = Callable[[AnalyticsExportJob, AnalyticsFilters, str], tuple[Path, int]]

EXPORT_HANDLERS: dict[str, ExportHandler] = {}


def export_report(export_job: AnalyticsExportJob) -> tuple[Path, int]:
    filters = parse_filters_payload(export_job.filters)
    handler = EXPORT_HANDLERS.get(export_job.export_type)
    if handler is None:
        raise ValueError(f"Unsupported export type: {export_job.export_type}")
    return handler(export_job, filters, export_job.file_format)


def _export_dir(export_job: AnalyticsExportJob) -> Path:
    relative = Path("exports") / "analytics" / str(export_job.id)
    absolute = Path(settings.MEDIA_ROOT) / relative
    absolute.mkdir(parents=True, exist_ok=True)
    return relative


def _export_tasks(
    export_job: AnalyticsExportJob,
    filters: AnalyticsFilters,
    file_format: str,
) -> tuple[Path, int]:
    rows = _task_rows(_tasks_queryset(filters))
    headers = [
        "id",
        "title",
        "column",
        "system_type",
        "week",
        "tags",
        "source",
        "due_at",
        "created_at",
        "closed_at",
        "archived_at",
        "evidence_url",
    ]
    return _write_output(export_job, file_format, "tasks", [("Tasks", headers, rows)])


def _export_archive(
    export_job: AnalyticsExportJob,
    filters: AnalyticsFilters,
    file_format: str,
) -> tuple[Path, int]:
    board = ColumnSettingsService.get_default_board()
    queryset = (
        selectors._apply_filters(
            Task.objects.filter(board=board, archived_at__isnull=False),
            filters,
            period=False,
        )
        .select_related("column", "week")
        .prefetch_related("tags")
    )
    headers = [
        "id",
        "title",
        "column",
        "system_type",
        "week",
        "tags",
        "source",
        "due_at",
        "created_at",
        "closed_at",
        "archived_at",
        "evidence_url",
        "completion_note",
    ]
    rows = [[*_task_row(task), task.completion_note] for task in queryset]
    return _write_output(
        export_job,
        file_format,
        "archive",
        [("Archive", headers, rows)],
    )


def _export_weekly_summary(
    export_job: AnalyticsExportJob,
    filters: AnalyticsFilters,
    file_format: str,
) -> tuple[Path, int]:
    trend = selectors.get_weekly_trend(filters)
    summary_headers = ["week", "created", "closed", "carried_over"]
    summary_rows = [
        [item["week"], item["created"], item["closed"], item["carried_over"]]
        for item in trend["items"]
    ]

    tag_breakdown = selectors.get_breakdown(filters, group_by="tag")
    tag_headers = ["tag", "count"]
    tag_rows = [[item["key"], item["count"]] for item in tag_breakdown["items"]]

    column_breakdown = selectors.get_breakdown(filters, group_by="column")
    column_headers = ["column", "system_type", "count"]
    column_rows = [
        [item["key"], item.get("system_type", ""), item["count"]]
        for item in column_breakdown["items"]
    ]

    stale_board = ColumnSettingsService.get_default_board()
    active_qs = selectors._apply_filters(
        Task.objects.filter(board=stale_board, archived_at__isnull=True),
        filters,
        period=False,
    )
    stale_qs = selectors._stale_tasks_queryset(active_qs)
    stale_headers = ["id", "title", "column", "days_in_column"]
    stale_rows = [
        [item["id"], item["title"], item["column"], item["days_in_column"]]
        for item in selectors._stale_task_items(stale_qs)
    ]

    sheets = [
        ("Summary", summary_headers, summary_rows),
        ("By Tags", tag_headers, tag_rows),
        ("By Columns", column_headers, column_rows),
        ("Stale Tasks", stale_headers, stale_rows),
    ]
    return _write_output(export_job, file_format, "weekly_summary", sheets)


def _export_tag_summary(
    export_job: AnalyticsExportJob,
    filters: AnalyticsFilters,
    file_format: str,
) -> tuple[Path, int]:
    breakdown = selectors.get_breakdown(filters, group_by="tag")
    headers = ["tag", "count"]
    rows = [[item["key"], item["count"]] for item in breakdown["items"]]
    return _write_output(
        export_job, file_format, "tag_summary", [("Tags", headers, rows)]
    )


def _export_notification_report(
    export_job: AnalyticsExportJob,
    filters: AnalyticsFilters,
    file_format: str,
) -> tuple[Path, int]:
    period_start, period_end = period_bounds(filters)
    jobs = selectors._filter_datetime(
        NotificationJob.objects.select_related("task"),
        "created_at",
        period_start,
        period_end,
    )
    headers = [
        "id",
        "task_id",
        "task_title",
        "status",
        "reason",
        "telegram_chat_id",
        "scheduled_at",
        "sent_at",
        "failed_at",
        "last_error",
    ]
    rows = [
        [
            job.id,
            job.task_id,
            job.task.title,
            job.status,
            job.reason,
            job.telegram_chat_id,
            _iso(job.scheduled_at),
            _iso(job.sent_at),
            _iso(job.failed_at),
            job.last_error,
        ]
        for job in jobs
    ]
    return _write_output(
        export_job,
        file_format,
        "notification_report",
        [("Notifications", headers, rows)],
    )


def _export_jobs_report(
    export_job: AnalyticsExportJob,
    filters: AnalyticsFilters,
    file_format: str,
) -> tuple[Path, int]:
    period_start, period_end = period_bounds(filters)
    jobs = selectors._filter_datetime(
        BackgroundJob.objects.all(), "created_at", period_start, period_end
    )
    headers = [
        "id",
        "job_type",
        "status",
        "attempts",
        "scheduled_at",
        "started_at",
        "finished_at",
        "last_error",
        "created_at",
    ]
    rows = [
        [
            job.id,
            job.job_type,
            job.status,
            job.attempts,
            _iso(job.scheduled_at),
            _iso(job.started_at),
            _iso(job.finished_at),
            job.last_error,
            _iso(job.created_at),
        ]
        for job in jobs
    ]
    return _write_output(
        export_job, file_format, "jobs_report", [("Jobs", headers, rows)]
    )


def _export_imports_report(
    export_job: AnalyticsExportJob,
    filters: AnalyticsFilters,
    file_format: str,
) -> tuple[Path, int]:
    period_start, period_end = period_bounds(filters)
    logs = selectors._filter_datetime(
        ImportLog.objects.all(), "created_at", period_start, period_end
    )
    headers = [
        "id",
        "filename",
        "status",
        "tasks_created",
        "source_label",
        "error_message",
        "started_at",
        "finished_at",
        "created_at",
    ]
    rows = [
        [
            log.id,
            log.filename,
            log.status,
            log.tasks_created,
            log.source_label,
            log.error_message,
            _iso(log.started_at),
            _iso(log.finished_at),
            _iso(log.created_at),
        ]
        for log in logs
    ]
    return _write_output(
        export_job, file_format, "imports_report", [("Imports", headers, rows)]
    )


def _tasks_queryset(filters: AnalyticsFilters):
    board = ColumnSettingsService.get_default_board()
    return selectors._apply_filters(
        Task.objects.filter(board=board), filters, period=False
    )


def _task_rows(queryset) -> list[list]:
    queryset = queryset.select_related("column", "week").prefetch_related("tags")
    return [_task_row(task) for task in queryset]


def _task_row(task: Task) -> list:
    week = ""
    if task.week is not None:
        week = f"{task.week.iso_year}-W{task.week.iso_week:02d}"
    tags = ", ".join(tag.name for tag in task.tags.all())
    return [
        task.id,
        task.title,
        task.column.name,
        task.column.system_type,
        week,
        tags,
        task.source,
        _iso(task.due_at),
        _iso(task.created_at),
        _iso(task.closed_at),
        _iso(task.archived_at),
        task.evidence_url,
    ]


def _write_output(
    export_job: AnalyticsExportJob,
    file_format: str,
    basename: str,
    sheets: list[tuple[str, list[str], list[list]]],
) -> tuple[Path, int]:
    export_dir = _export_dir(export_job)
    if file_format == FileFormat.CSV:
        _name, headers, rows = sheets[0]
        filename = f"{basename}.csv"
        path = Path(settings.MEDIA_ROOT) / export_dir / filename
        _write_csv(path, headers, rows)
    else:
        filename = f"{basename}.xlsx"
        path = Path(settings.MEDIA_ROOT) / export_dir / filename
        _write_xlsx(path, sheets)

    row_count = sum(len(sheet[2]) for sheet in sheets)
    return export_dir / filename, row_count


def _write_csv(path: Path, headers: list[str], rows: list[list]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def _write_xlsx(path: Path, sheets: list[tuple[str, list[str], list[list]]]) -> None:
    workbook = Workbook()
    active = workbook.active
    if active is not None:
        workbook.remove(active)

    for index, (title, headers, rows) in enumerate(sheets):
        worksheet = workbook.create_sheet(title=title[:31], index=index)
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for column_index, header in enumerate(headers, start=1):
            column = get_column_letter(column_index)
            max_length = len(str(header))
            for row in rows:
                value = row[column_index - 1] if column_index - 1 < len(row) else ""
                max_length = max(max_length, len(str(value)))
            worksheet.column_dimensions[column].width = min(max_length + 2, 48)

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

    workbook.save(path)


def _iso(value: datetime | None) -> str:
    if value is None:
        return ""
    if timezone.is_naive(value):
        value = timezone.make_aware(value)
    return value.isoformat()


EXPORT_HANDLERS = {
    ExportType.TASKS: _export_tasks,
    ExportType.ARCHIVE: _export_archive,
    ExportType.WEEKLY_SUMMARY: _export_weekly_summary,
    ExportType.TAG_SUMMARY: _export_tag_summary,
    ExportType.NOTIFICATION_REPORT: _export_notification_report,
    ExportType.JOBS_REPORT: _export_jobs_report,
    ExportType.IMPORTS_REPORT: _export_imports_report,
}
