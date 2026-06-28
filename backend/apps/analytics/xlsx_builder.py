from __future__ import annotations

import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Literal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from apps.analytics.export_i18n import et

FONT_NAME = "Calibri"
MARGIN_COL = 1
CONTENT_START_COL = 2
MIN_CONTENT_COLS = 10
MAX_CONTENT_COLS = 14
CONTENT_TARGET_WIDTH = 98.0

COLOR_PRIMARY = "5B6EE8"
COLOR_ACCENT = "4A7FD6"
COLOR_WARN = "F59E0B"
COLOR_VIOLET = "8B5CF6"
COLOR_INDIGO = "6366F1"
COLOR_SLATE = "64748B"
COLOR_TEXT = "1E293B"
COLOR_TEXT_MUTED = "64748B"
COLOR_META_LABEL = "475569"
COLOR_HEADER = "1E40AF"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_TITLE_BAR = "1E3A8A"
COLOR_SURFACE = "F8FAFC"
COLOR_SURFACE_SOFT = "F1F5F9"
COLOR_BORDER = "E2E8F0"
COLOR_CARD_EDGE = "CBD5E1"
COLOR_KPI_VALUE = "1E40AF"

TITLE_BAR_FILL = PatternFill("solid", fgColor=COLOR_TITLE_BAR)
HEADER_FILL = PatternFill("solid", fgColor=COLOR_HEADER)
META_FILL = PatternFill("solid", fgColor=COLOR_SURFACE)
KPI_LABEL_FILL = PatternFill("solid", fgColor=COLOR_SURFACE_SOFT)
KPI_VALUE_FILL = PatternFill("solid", fgColor="FFFFFF")
ZEBRA_FILL = PatternFill("solid", fgColor=COLOR_SURFACE)

TITLE_FONT = Font(name=FONT_NAME, color=COLOR_HEADER_TEXT, bold=True, size=17)
META_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=10, color=COLOR_META_LABEL)
META_FONT = Font(name=FONT_NAME, size=10, color=COLOR_TEXT_MUTED)
KPI_LABEL_FONT = Font(name=FONT_NAME, color=COLOR_TEXT_MUTED, size=9, bold=True)
KPI_VALUE_FONT = Font(name=FONT_NAME, color=COLOR_KPI_VALUE, bold=True, size=15)
HEADER_FONT = Font(name=FONT_NAME, color=COLOR_HEADER_TEXT, bold=True, size=11)
BODY_FONT = Font(name=FONT_NAME, size=10, color=COLOR_TEXT)
BODY_NUM_FONT = Font(name=FONT_NAME, size=10, color=COLOR_TEXT, bold=False)
EMPTY_FONT = Font(name=FONT_NAME, size=11, color=COLOR_TEXT_MUTED, italic=True)

SIDE_EDGE = Side(style="medium", color=COLOR_CARD_EDGE)
THIN_SIDE = Side(style="thin", color=COLOR_BORDER)

THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)
META_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_CENTER_NOWRAP = Alignment(horizontal="center", vertical="center", wrap_text=False)
ALIGN_LEFT_CENTER = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_LEFT_CENTER_INDENT = Alignment(
    horizontal="left", vertical="center", wrap_text=True, indent=1
)

META_INNER_MIN_COLS = 6
META_INNER_MAX_COLS = 9
DATA_TABLE_MIN_COLS = 5
DATA_TABLE_MIN_ROWS = 12


@dataclass
class ColumnProfile:
    min_width: float
    max_width: float
    alignment: Alignment


COLUMN_PROFILES: dict[str, ColumnProfile] = {
    "export.col.id": ColumnProfile(8, 12, ALIGN_CENTER_NOWRAP),
    "export.col.title": ColumnProfile(28, 50, ALIGN_CENTER),
    "export.col.column": ColumnProfile(14, 24, ALIGN_CENTER),
    "export.col.systemType": ColumnProfile(12, 20, ALIGN_CENTER),
    "export.col.week": ColumnProfile(12, 18, ALIGN_CENTER_NOWRAP),
    "export.col.tags": ColumnProfile(22, 42, ALIGN_CENTER),
    "export.col.source": ColumnProfile(12, 20, ALIGN_CENTER),
    "export.col.dueAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.createdAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.closedAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.archivedAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.completionNote": ColumnProfile(24, 48, ALIGN_CENTER),
    "export.col.daysInColumn": ColumnProfile(12, 18, ALIGN_CENTER_NOWRAP),
    "export.col.tag": ColumnProfile(18, 36, ALIGN_CENTER),
    "export.col.count": ColumnProfile(10, 16, ALIGN_CENTER_NOWRAP),
    "export.col.created": ColumnProfile(10, 14, ALIGN_CENTER_NOWRAP),
    "export.col.closed": ColumnProfile(10, 14, ALIGN_CENTER_NOWRAP),
    "export.col.carriedOver": ColumnProfile(12, 16, ALIGN_CENTER_NOWRAP),
    "export.col.taskId": ColumnProfile(8, 12, ALIGN_CENTER_NOWRAP),
    "export.col.taskTitle": ColumnProfile(28, 50, ALIGN_CENTER),
    "export.col.status": ColumnProfile(12, 20, ALIGN_CENTER),
    "export.col.reason": ColumnProfile(18, 36, ALIGN_CENTER),
    "export.col.telegramChat": ColumnProfile(16, 28, ALIGN_CENTER),
    "export.col.scheduledAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.sentAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.failedAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.lastError": ColumnProfile(24, 48, ALIGN_CENTER),
    "export.col.jobType": ColumnProfile(16, 28, ALIGN_CENTER),
    "export.col.attempts": ColumnProfile(10, 14, ALIGN_CENTER_NOWRAP),
    "export.col.startedAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.finishedAt": ColumnProfile(19, 22, ALIGN_CENTER_NOWRAP),
    "export.col.filename": ColumnProfile(20, 40, ALIGN_CENTER),
    "export.col.tasksCreated": ColumnProfile(12, 16, ALIGN_CENTER_NOWRAP),
    "export.col.sourceLabel": ColumnProfile(16, 28, ALIGN_CENTER),
    "export.col.errorMessage": ColumnProfile(24, 48, ALIGN_CENTER),
}
DEFAULT_COLUMN_PROFILE = ColumnProfile(12, 28, ALIGN_CENTER)

PIE_SLICE_COLORS = (
    COLOR_PRIMARY,
    COLOR_ACCENT,
    COLOR_VIOLET,
    COLOR_INDIGO,
    COLOR_WARN,
    COLOR_SLATE,
    "0891B2",
)

ChartKind = Literal["weekly_combo", "pie", "bar_col"]


@dataclass
class ExportMeta:
    locale: str
    report_title: str
    scheme_name: str = ""
    period_label: str = ""
    generated_at: datetime | None = None


@dataclass
class SheetChartSpec:
    kind: ChartKind
    title_key: str


@dataclass
class ColSpan:
    start: int
    end: int

    @property
    def anchor(self) -> int:
        return self.start


@dataclass
class TableBounds:
    header_row: int
    data_start_row: int
    data_end_row: int
    column_count: int
    start_col: int
    end_col: int
    col_spans: list[ColSpan]
    compact: bool
    header_keys: list[str]


@dataclass
class SheetLayout:
    start_col: int
    span: int

    @property
    def end_col(self) -> int:
        return self.start_col + self.span - 1


class AnalyticsXlsxBuilder:
    def __init__(self, meta: ExportMeta) -> None:
        self.meta = meta
        self.workbook = Workbook()
        if self.workbook.active is not None:
            self.workbook.remove(self.workbook.active)

    def add_table_sheet(
        self,
        *,
        sheet_key: str,
        header_keys: list[str],
        rows: list[list],
        kpis: list[tuple[str, object]] | None = None,
        charts: list[SheetChartSpec] | None = None,
    ) -> None:
        title = et(sheet_key, locale=self.meta.locale)
        worksheet = self.workbook.create_sheet(title=title[:31])
        self._prepare_worksheet(worksheet)

        headers = [et(key, locale=self.meta.locale) for key in header_keys]
        kpi_count = len(kpis) if kpis else 0
        compact_table = self._is_compact_table(header_keys, rows)
        layout = self._compute_layout(
            table_cols=len(headers),
            kpi_count=kpi_count,
            compact_table=compact_table,
        )

        next_row = self._write_title_band(worksheet, title, layout)
        next_row = self._write_meta_block(worksheet, next_row, layout)
        if kpis:
            next_row = self._section_gap(worksheet, next_row, layout)
            next_row = self._write_kpi_block(worksheet, next_row, kpis, layout)
        next_row = self._section_gap(worksheet, next_row, layout)
        bounds = self._write_table(
            worksheet,
            next_row,
            headers,
            header_keys,
            rows,
            layout,
            compact=compact_table,
        )

        last_row = (
            bounds.data_end_row
            if rows or bounds.data_end_row >= bounds.data_start_row
            else bounds.header_row
        )
        if charts and rows:
            last_row = bounds.data_end_row + 2
            for chart_spec in charts:
                last_row = self._add_chart(
                    worksheet, bounds, chart_spec, layout, last_row
                )

        self._apply_sheet_column_widths(worksheet, layout, bounds, headers, rows)
        if not compact_table:
            self._autofit_row_heights(worksheet, bounds, rows)
        self._apply_card_frame(worksheet, layout, first_row=1, last_row=last_row)
        self._finalize_sheet_viewport(worksheet, layout, compact_table=compact_table)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        if self.workbook.worksheets:
            self.workbook.active = self.workbook.worksheets[0]
        self.workbook.save(path)

    def _compute_layout(
        self,
        *,
        table_cols: int,
        kpi_count: int,
        compact_table: bool,
    ) -> SheetLayout:
        if compact_table:
            needed = max(table_cols, kpi_count, MIN_CONTENT_COLS)
            span = min(needed, MAX_CONTENT_COLS)
        else:
            span = table_cols
        return SheetLayout(start_col=CONTENT_START_COL, span=span)

    @staticmethod
    def _is_compact_table(header_keys: list[str], rows: list[list]) -> bool:
        if len(header_keys) >= DATA_TABLE_MIN_COLS:
            return False
        if len(rows) >= DATA_TABLE_MIN_ROWS:
            return False
        long_values = any(len(str(cell)) > 24 for row in rows[:40] for cell in row)
        return not (long_values and len(header_keys) >= 4)

    @staticmethod
    def _column_profile(key: str) -> ColumnProfile:
        return COLUMN_PROFILES.get(key, DEFAULT_COLUMN_PROFILE)

    @staticmethod
    def _cell_display_length(value: object) -> int:
        text = str(value) if value is not None else ""
        return max((len(part) for part in text.splitlines()), default=0)

    def _column_content_width(
        self,
        header: str,
        header_key: str,
        rows: list[list],
        column_index: int,
    ) -> float:
        profile = self._column_profile(header_key)
        max_length = max(self._cell_display_length(header), profile.min_width - 2)
        for row in rows[:120]:
            if column_index < len(row):
                max_length = max(
                    max_length, self._cell_display_length(row[column_index])
                )
        return min(max(max_length + 2, profile.min_width), profile.max_width)

    @staticmethod
    def _distribute_spans(layout: SheetLayout, block_count: int) -> list[ColSpan]:
        if block_count <= 0:
            return []
        base = layout.span // block_count
        remainder = layout.span % block_count
        spans: list[ColSpan] = []
        col = layout.start_col
        for index in range(block_count):
            width = base + (1 if index < remainder else 0)
            spans.append(ColSpan(start=col, end=col + width - 1))
            col += width
        return spans

    def _prepare_worksheet(self, worksheet: Worksheet) -> None:
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_format.defaultRowHeight = 15
        worksheet.sheet_view.zoomScale = 110
        worksheet.sheet_view.zoomScaleNormal = 110
        worksheet.print_options.horizontalCentered = True
        worksheet.page_margins.left = 0.35
        worksheet.page_margins.right = 0.35
        worksheet.page_margins.top = 0.45
        worksheet.page_margins.bottom = 0.45
        worksheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

    def _section_gap(self, worksheet: Worksheet, row: int, layout: SheetLayout) -> int:
        worksheet.row_dimensions[row].height = 6
        self._paint_row_edge(worksheet, row, layout)
        return row + 1

    def _paint_row_edge(
        self, worksheet: Worksheet, row: int, layout: SheetLayout
    ) -> None:
        for col in (layout.start_col, layout.end_col):
            cell = worksheet.cell(row=row, column=col)
            cell.border = Border(
                left=SIDE_EDGE if col == layout.start_col else cell.border.left,
                right=SIDE_EDGE if col == layout.end_col else cell.border.right,
                top=cell.border.top,
                bottom=cell.border.bottom,
            )

    def _merge_write(
        self,
        worksheet: Worksheet,
        row: int,
        span: ColSpan,
        *,
        value,
        font: Font,
        fill: PatternFill | None = None,
        alignment: Alignment = ALIGN_CENTER,
        border: Border | None = THIN_BORDER,
    ) -> None:
        if span.start < span.end:
            worksheet.merge_cells(
                start_row=row,
                start_column=span.start,
                end_row=row,
                end_column=span.end,
            )
        cell = worksheet.cell(row=row, column=span.start, value=value)
        cell.font = font
        cell.alignment = alignment
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border

    def _write_cell(
        self,
        worksheet: Worksheet,
        row: int,
        col: int,
        *,
        value,
        font: Font,
        fill: PatternFill | None = None,
        alignment: Alignment = ALIGN_CENTER,
        border: Border | None = THIN_BORDER,
    ) -> None:
        cell = worksheet.cell(row=row, column=col, value=value)
        cell.font = font
        cell.alignment = alignment
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border

    def _write_title_band(
        self, worksheet: Worksheet, sheet_title: str, layout: SheetLayout
    ) -> int:
        span = ColSpan(start=layout.start_col, end=layout.end_col)
        self._merge_write(
            worksheet,
            1,
            span,
            value=sheet_title,
            font=TITLE_FONT,
            fill=TITLE_BAR_FILL,
            alignment=ALIGN_CENTER,
            border=Border(
                left=SIDE_EDGE,
                right=SIDE_EDGE,
                top=SIDE_EDGE,
                bottom=Side(style="thin", color=COLOR_TITLE_BAR),
            ),
        )
        worksheet.row_dimensions[1].height = 40
        return 2

    def _meta_inner_spans(self, layout: SheetLayout) -> tuple[ColSpan, ColSpan]:
        inner_span = min(
            max(META_INNER_MIN_COLS, layout.span // 2 + 2),
            layout.span,
            META_INNER_MAX_COLS,
        )
        inner_start = layout.start_col + (layout.span - inner_span) // 2
        inner_end = inner_start + inner_span - 1
        label_cols = max(2, inner_span * 2 // 5)
        label_end = inner_start + label_cols - 1
        return (
            ColSpan(start=inner_start, end=label_end),
            ColSpan(start=label_end + 1, end=inner_end),
        )

    def _write_meta_block(
        self, worksheet: Worksheet, start_row: int, layout: SheetLayout
    ) -> int:
        generated = self.meta.generated_at or timezone.now()
        generated_local = timezone.localtime(generated)
        meta_lines = [
            (et("export.reportTitle", locale=self.meta.locale), self.meta.report_title),
        ]
        if self.meta.scheme_name:
            meta_lines.append(
                (et("export.scheme", locale=self.meta.locale), self.meta.scheme_name),
            )
        if self.meta.period_label:
            meta_lines.append(
                (et("export.period", locale=self.meta.locale), self.meta.period_label),
            )
        meta_lines.append(
            (
                et("export.generatedAt", locale=self.meta.locale),
                generated_local.strftime("%d.%m.%Y, %H:%M"),
            ),
        )

        label_span, value_span = self._meta_inner_spans(layout)

        row = start_row
        for label, value in meta_lines:
            self._merge_write(
                worksheet,
                row,
                label_span,
                value=label,
                font=META_LABEL_FONT,
                fill=META_FILL,
                alignment=ALIGN_LEFT_CENTER_INDENT,
                border=META_BORDER,
            )
            self._merge_write(
                worksheet,
                row,
                value_span,
                value=value,
                font=META_FONT,
                fill=META_FILL,
                alignment=ALIGN_LEFT_CENTER_INDENT,
                border=META_BORDER,
            )
            worksheet.row_dimensions[row].height = 22
            row += 1
        return row

    def _write_kpi_block(
        self,
        worksheet: Worksheet,
        start_row: int,
        kpis: list[tuple[str, object]],
        layout: SheetLayout,
    ) -> int:
        spans = self._distribute_spans(layout, len(kpis))
        for index, (label_key, value) in enumerate(kpis):
            label_text = et(label_key, locale=self.meta.locale)
            self._merge_write(
                worksheet,
                start_row,
                spans[index],
                value=label_text,
                font=KPI_LABEL_FONT,
                fill=KPI_LABEL_FILL,
            )
            self._merge_write(
                worksheet,
                start_row + 1,
                spans[index],
                value=value,
                font=KPI_VALUE_FONT,
                fill=KPI_VALUE_FILL,
            )

        worksheet.row_dimensions[start_row].height = 24
        worksheet.row_dimensions[start_row + 1].height = 32
        return start_row + 2

    def _write_table(
        self,
        worksheet: Worksheet,
        start_row: int,
        headers: list[str],
        header_keys: list[str],
        rows: list[list],
        layout: SheetLayout,
        *,
        compact: bool,
    ) -> TableBounds:
        if compact:
            col_spans = self._distribute_spans(layout, len(headers))
        else:
            col_spans = [
                ColSpan(start=layout.start_col + index, end=layout.start_col + index)
                for index in range(len(headers))
            ]

        worksheet.row_dimensions[start_row].height = 30 if compact else 34

        for column_index, header in enumerate(headers):
            if compact:
                self._merge_write(
                    worksheet,
                    start_row,
                    col_spans[column_index],
                    value=header,
                    font=HEADER_FONT,
                    fill=HEADER_FILL,
                    alignment=ALIGN_CENTER,
                )
            else:
                self._write_cell(
                    worksheet,
                    start_row,
                    col_spans[column_index].start,
                    value=header,
                    font=HEADER_FONT,
                    fill=HEADER_FILL,
                    alignment=ALIGN_CENTER,
                )

        data_start_row = start_row + 1
        if not rows and headers:
            empty_span = ColSpan(start=layout.start_col, end=layout.end_col)
            self._merge_write(
                worksheet,
                data_start_row,
                empty_span,
                value=et("export.table.noData", locale=self.meta.locale),
                font=EMPTY_FONT,
                fill=META_FILL,
                alignment=ALIGN_CENTER,
            )
            worksheet.row_dimensions[data_start_row].height = 36
            data_end_row = data_start_row
        else:
            for row_index, row in enumerate(rows, start=data_start_row):
                fill = ZEBRA_FILL if row_index % 2 == 0 else None
                for column_index, value in enumerate(row):
                    profile = self._column_profile(header_keys[column_index])
                    cell_font = BODY_FONT if column_index == 0 else BODY_NUM_FONT
                    if compact:
                        self._merge_write(
                            worksheet,
                            row_index,
                            col_spans[column_index],
                            value=value,
                            font=cell_font,
                            fill=fill,
                            alignment=profile.alignment,
                        )
                    else:
                        self._write_cell(
                            worksheet,
                            row_index,
                            col_spans[column_index].start,
                            value=value,
                            font=cell_font,
                            fill=fill,
                            alignment=profile.alignment,
                        )

            data_end_row = (
                data_start_row - 1 if not rows else data_start_row + len(rows) - 1
            )

        if headers and rows:
            worksheet.freeze_panes = (
                f"{get_column_letter(layout.start_col)}{start_row + 1}"
            )
            worksheet.auto_filter.ref = (
                f"{get_column_letter(layout.start_col)}{start_row}:"
                f"{get_column_letter(layout.end_col)}{data_end_row}"
            )

        return TableBounds(
            header_row=start_row,
            data_start_row=data_start_row,
            data_end_row=data_end_row,
            column_count=len(headers),
            start_col=layout.start_col,
            end_col=layout.end_col,
            col_spans=col_spans,
            compact=compact,
            header_keys=header_keys,
        )

    def _apply_sheet_column_widths(
        self,
        worksheet: Worksheet,
        layout: SheetLayout,
        bounds: TableBounds,
        headers: list[str],
        rows: list[list],
    ) -> None:
        if bounds.compact:
            width = round(CONTENT_TARGET_WIDTH / layout.span, 2)
            for col in range(layout.start_col, layout.end_col + 1):
                worksheet.column_dimensions[get_column_letter(col)].width = width
            return

        for column_index, header_key in enumerate(bounds.header_keys):
            col = layout.start_col + column_index
            width = self._column_content_width(
                headers[column_index],
                header_key,
                rows,
                column_index,
            )
            worksheet.column_dimensions[get_column_letter(col)].width = round(width, 2)

    def _autofit_row_heights(
        self, worksheet: Worksheet, bounds: TableBounds, rows: list[list]
    ) -> None:
        for row_index, row in enumerate(rows, start=bounds.data_start_row):
            max_lines = 1
            for column_index, value in enumerate(row):
                profile = self._column_profile(bounds.header_keys[column_index])
                col_letter = get_column_letter(bounds.start_col + column_index)
                col_width = (
                    worksheet.column_dimensions[col_letter].width or profile.min_width
                )
                text = str(value) if value is not None else ""
                if profile.alignment.wrap_text:
                    explicit_lines = text.count("\n") + 1
                    wrapped_lines = max(
                        1,
                        math.ceil(
                            self._cell_display_length(text) / max(col_width - 1, 8)
                        ),
                    )
                    max_lines = max(max_lines, explicit_lines, wrapped_lines)
            worksheet.row_dimensions[row_index].height = min(16 * max_lines + 6, 96)

        header_lines = 1
        for column_index, header in enumerate(bounds.header_keys):
            profile = self._column_profile(header)
            col_letter = get_column_letter(bounds.start_col + column_index)
            col_width = (
                worksheet.column_dimensions[col_letter].width or profile.min_width
            )
            header_text = worksheet.cell(
                bounds.header_row, bounds.start_col + column_index
            ).value
            header_text = str(header_text) if header_text is not None else ""
            header_lines = max(
                header_lines,
                math.ceil(len(header_text) / max(col_width - 1, 8)),
            )
        worksheet.row_dimensions[bounds.header_row].height = min(
            16 * header_lines + 10, 48
        )

    def _apply_card_frame(
        self,
        worksheet: Worksheet,
        layout: SheetLayout,
        *,
        first_row: int,
        last_row: int,
    ) -> None:
        for row in range(first_row, last_row + 1):
            for col in (layout.start_col, layout.end_col):
                cell = worksheet.cell(row=row, column=col)
                cell.border = Border(
                    left=SIDE_EDGE if col == layout.start_col else cell.border.left,
                    right=SIDE_EDGE if col == layout.end_col else cell.border.right,
                    top=cell.border.top,
                    bottom=cell.border.bottom,
                )

        for col in range(layout.start_col, layout.end_col + 1):
            top = worksheet.cell(row=first_row, column=col)
            bottom = worksheet.cell(row=last_row, column=col)
            top.border = Border(
                left=top.border.left,
                right=top.border.right,
                top=SIDE_EDGE,
                bottom=top.border.bottom,
            )
            bottom.border = Border(
                left=bottom.border.left,
                right=bottom.border.right,
                top=bottom.border.top,
                bottom=SIDE_EDGE,
            )

    def _finalize_sheet_viewport(
        self,
        worksheet: Worksheet,
        layout: SheetLayout,
        *,
        compact_table: bool,
    ) -> None:
        worksheet.column_dimensions[get_column_letter(MARGIN_COL)].width = 1.1

        right_margin_col = layout.end_col + 1
        worksheet.column_dimensions[get_column_letter(right_margin_col)].width = 1.1

        hide_from = right_margin_col + 1
        for col_idx in range(hide_from, hide_from + 60):
            letter = get_column_letter(col_idx)
            dim = worksheet.column_dimensions[letter]
            dim.hidden = True
            dim.width = 0.1

        if compact_table:
            worksheet.page_setup.orientation = (
                "landscape" if layout.span >= 8 else "portrait"
            )
        else:
            worksheet.page_setup.orientation = "landscape"
            worksheet.page_setup.fitToWidth = 0
            worksheet.sheet_properties.pageSetUpPr = PageSetupProperties(
                fitToPage=False
            )

    def _chart_size(self, layout: SheetLayout, kind: ChartKind) -> tuple[float, float]:
        if kind == "pie":
            return min(20.0, 3.5 + layout.span * 1.5), 14.0
        return min(26.0, 4.5 + layout.span * 1.85), 13.5

    def _add_chart(
        self,
        worksheet: Worksheet,
        bounds: TableBounds,
        spec: SheetChartSpec,
        layout: SheetLayout,
        anchor_row: int,
    ) -> int:
        title = et(spec.title_key, locale=self.meta.locale)
        if spec.kind == "weekly_combo":
            chart = self._build_weekly_combo_chart(worksheet, bounds, title)
        elif spec.kind == "pie":
            chart = self._build_pie_chart(worksheet, bounds, title)
        else:
            chart = self._build_bar_chart(worksheet, bounds, title)

        width, height = self._chart_size(layout, spec.kind)
        chart.width = width
        chart.height = height
        anchor_col = get_column_letter(layout.start_col)
        worksheet.add_chart(chart, f"{anchor_col}{anchor_row}")
        chart_rows = int(height * 2.2)
        return anchor_row + chart_rows

    def _series_ref(
        self,
        worksheet: Worksheet,
        bounds: TableBounds,
        column_index: int,
    ) -> Reference:
        span = bounds.col_spans[column_index]
        return Reference(
            worksheet,
            min_col=span.anchor,
            min_row=bounds.header_row,
            max_row=bounds.data_end_row,
        )

    def _category_ref(self, worksheet: Worksheet, bounds: TableBounds) -> Reference:
        span = bounds.col_spans[0]
        return Reference(
            worksheet,
            min_col=span.anchor,
            min_row=bounds.data_start_row,
            max_row=bounds.data_end_row,
        )

    def _build_weekly_combo_chart(
        self,
        worksheet: Worksheet,
        bounds: TableBounds,
        title: str,
    ) -> BarChart:
        bar = BarChart()
        bar.type = "col"
        bar.grouping = "clustered"
        bar.style = 2
        bar.title = title
        bar.legend.position = "b"
        bar.gapWidth = 80
        bar.varyColors = False

        categories = self._category_ref(worksheet, bounds)
        bar_data = Reference(
            worksheet,
            min_col=bounds.col_spans[1].anchor,
            min_row=bounds.header_row,
            max_col=bounds.col_spans[2].anchor,
            max_row=bounds.data_end_row,
        )
        bar.add_data(bar_data, titles_from_data=True)
        bar.set_categories(categories)

        bar_colors = [COLOR_PRIMARY, COLOR_ACCENT]
        for index, series in enumerate(bar.series):
            series.graphicalProperties.solidFill = bar_colors[index % len(bar_colors)]

        line = LineChart()
        line.style = 2
        line_data = Reference(
            worksheet,
            min_col=bounds.col_spans[3].anchor,
            min_row=bounds.header_row,
            max_row=bounds.data_end_row,
        )
        line.add_data(line_data, titles_from_data=True)
        line.set_categories(categories)
        if line.series:
            series = line.series[0]
            series.graphicalProperties.line.solidFill = COLOR_WARN
            series.graphicalProperties.line.width = 24000
            series.marker.symbol = "circle"
            series.marker.size = 8
            series.marker.graphicalProperties.solidFill = COLOR_WARN
            series.marker.graphicalProperties.line.solidFill = COLOR_WARN

        line.y_axis.axId = 200
        bar.y_axis.crossAx = 200
        bar += line
        return bar

    def _build_pie_chart(
        self,
        worksheet: Worksheet,
        bounds: TableBounds,
        title: str,
    ) -> PieChart:
        pie = PieChart()
        pie.title = title
        pie.style = 2

        labels = self._category_ref(worksheet, bounds)
        data = self._series_ref(worksheet, bounds, 1)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        if pie.series:
            self._apply_pie_colors(
                pie.series[0], bounds.data_end_row - bounds.data_start_row + 1
            )
        return pie

    def _apply_pie_colors(self, series, slice_count: int) -> None:
        series.dPt = []
        for index in range(slice_count):
            point = DataPoint(idx=index)
            point.graphicalProperties.solidFill = PIE_SLICE_COLORS[
                index % len(PIE_SLICE_COLORS)
            ]
            series.dPt.append(point)

    def _build_bar_chart(
        self,
        worksheet: Worksheet,
        bounds: TableBounds,
        title: str,
    ) -> BarChart:
        chart = BarChart()
        chart.type = "col"
        chart.style = 2
        chart.title = title
        chart.legend = None
        chart.gapWidth = 70
        chart.varyColors = False

        categories = self._category_ref(worksheet, bounds)
        data = self._series_ref(worksheet, bounds, bounds.column_count - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        if chart.series:
            chart.series[0].graphicalProperties.solidFill = COLOR_PRIMARY
        return chart
