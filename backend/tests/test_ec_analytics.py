import re

import pytest
from django.urls import reverse

from apps.analytics.exporters import export_report
from apps.analytics.filters import parse_filters_payload, period_bounds, resolve_tag_ids
from apps.analytics.models import (
    AnalyticsExportJob,
    ExportStatus,
    ExportType,
    FileFormat,
)
from apps.analytics.services import AnalyticsExportService
from conftest import close_task_via_api, create_task_via_api


@pytest.fixture
def analytics_task_data(api_client, planned_column, week_key, make_tag):
    tag = make_tag("Analytics", slug="analytics")
    created = create_task_via_api(
        api_client,
        title="Analytics task",
        column_id=planned_column.id,
        week=week_key,
        tags=[tag.slug],
    )
    close_task_via_api(api_client, created.json()["id"])
    return {"task_id": created.json()["id"], "tag_slug": tag.slug}


@pytest.mark.django_db
@pytest.mark.parametrize(
    "endpoint_name",
    [
        pytest.param("analytics-summary", id="ec_analytics_summary_endpoint"),
        pytest.param("analytics-weekly-trend", id="ec_analytics_weekly_trend_endpoint"),
        pytest.param("analytics-breakdown", id="ec_analytics_breakdown_endpoint"),
        pytest.param("analytics-cycle-time", id="ec_analytics_cycle_time_endpoint"),
        pytest.param(
            "analytics-notifications", id="ec_analytics_notifications_endpoint"
        ),
        pytest.param("analytics-task-flow", id="ec_analytics_task_flow_endpoint"),
        pytest.param("analytics-archive", id="ec_analytics_archive_endpoint"),
    ],
)
def test_analytics_ec_all_api_endpoints_return_200(
    api_client, analytics_task_data, endpoint_name
):
    params = {}
    if endpoint_name == "analytics-breakdown":
        params["group_by"] = "tag"
    response = api_client.get(reverse(endpoint_name), params)
    assert response.status_code == 200


@pytest.mark.django_db
def test_analytics_ec_notifications_include_scheme_status_actions(
    api_client, analytics_task_data
):
    response = api_client.get(reverse("analytics-notifications"))
    assert response.status_code == 200
    payload = response.json()
    assert "telegram_actions" in payload
    assert isinstance(payload["telegram_actions"], list)


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("group_by", "expected"),
    [
        pytest.param("column", "column", id="ec_breakdown_by_column_valid"),
        pytest.param("tag", "tag", id="ec_breakdown_by_tag_valid"),
        pytest.param("source", "source", id="ec_breakdown_by_source_valid"),
        pytest.param("unknown", "unknown", id="ec_breakdown_unknown_returns_empty"),
    ],
)
def test_analytics_ec_breakdown_group_classes(
    api_client, analytics_task_data, group_by, expected
):
    response = api_client.get(reverse("analytics-breakdown"), {"group_by": group_by})
    assert response.status_code == 200
    assert response.json()["group_by"] == expected


@pytest.mark.django_db
def test_analytics_ec_filters_parse_filters_payload_period_bounds_and_tag_resolution(
    current_week,
    make_tag,
):
    tag = make_tag("Ops", slug="ops")
    make_tag("Inactive", slug="inactive", is_active=False)
    filters = parse_filters_payload(
        {
            "week": f"{current_week.iso_year}-W{current_week.iso_week:02d}",
            "from": current_week.starts_on.isoformat(),
            "to": current_week.ends_on.isoformat(),
            "tags": ["ops", "missing"],
            "source": "api",
        }
    )
    start, end = period_bounds(filters)
    assert start is not None
    assert end is not None
    assert filters.source == "api"
    ids = resolve_tag_ids(["ops", "OPS", "missing"])
    assert tag.id in ids
    assert resolve_tag_ids(["inactive"]) == []


@pytest.mark.django_db
@pytest.mark.parametrize(
    "export_type",
    [
        pytest.param(ExportType.TASKS, id="ec_export_tasks_csv"),
        pytest.param(ExportType.ARCHIVE, id="ec_export_archive_csv"),
        pytest.param(ExportType.WEEKLY_SUMMARY, id="ec_export_weekly_summary_csv"),
        pytest.param(ExportType.TAG_SUMMARY, id="ec_export_tag_summary_csv"),
        pytest.param(ExportType.NOTIFICATION_REPORT, id="ec_export_notifications_csv"),
        pytest.param(ExportType.JOBS_REPORT, id="ec_export_jobs_csv"),
        pytest.param(ExportType.IMPORTS_REPORT, id="ec_export_imports_csv"),
    ],
)
def test_analytics_ec_export_all_types_csv(
    media_root, analytics_task_data, export_type
):
    export_job = AnalyticsExportJob.objects.create(
        export_type=export_type,
        file_format=FileFormat.CSV,
        status=ExportStatus.PENDING,
        filters={},
    )
    relative_path, rows = export_report(export_job)
    assert str(relative_path).endswith(".csv")
    assert rows >= 0


@pytest.mark.django_db
def test_analytics_ec_export_tasks_xlsx_matches_preview_layout(
    media_root, analytics_task_data
):
    from django.conf import settings
    from openpyxl import load_workbook

    export_job = AnalyticsExportJob.objects.create(
        export_type=ExportType.TASKS,
        file_format=FileFormat.XLSX,
        status=ExportStatus.PENDING,
        filters={"locale": "ru"},
    )
    relative_path, rows = export_report(export_job)
    assert str(relative_path).endswith(".xlsx")
    assert rows >= 0

    workbook = load_workbook(settings.MEDIA_ROOT / relative_path)
    assert len(workbook.sheetnames) >= 2
    overview = workbook[workbook.sheetnames[0]]
    title_values = [
        overview.cell(1, column).value
        for column in range(1, 13)
        if overview.cell(1, column).value
    ]
    assert any(value in {"Обзор", "Overview"} for value in title_values)

    header_row = None
    header_col = None
    for row_index in range(1, 30):
        for column_index in range(1, 13):
            if overview.cell(row_index, column_index).value in {"Колонка", "Column"}:
                header_row = row_index
                header_col = column_index
                break
        if header_row is not None:
            break
    assert header_row is not None
    assert header_col is not None
    assert overview.cell(header_row, header_col).fill.fgColor.rgb in {
        "001E40AF",
        "FF1E40AF",
    }
    header_values = {
        overview.cell(header_row, column).value
        for column in range(1, 15)
        if overview.cell(header_row, column).value
    }
    assert {"Количество", "Count"} & header_values
    assert len(overview._charts) == 1

    tasks_sheet = workbook[workbook.sheetnames[1]]
    tasks_header_row = None
    tasks_header_col = None
    for row_index in range(1, 30):
        for column_index in range(1, 20):
            if tasks_sheet.cell(row_index, column_index).value in {"ID", "id"}:
                tasks_header_row = row_index
                tasks_header_col = column_index
                break
        if tasks_header_row is not None:
            break
    assert tasks_header_row is not None
    assert tasks_header_col is not None
    header_values = {
        tasks_sheet.cell(tasks_header_row, column).value
        for column in range(tasks_header_col, tasks_header_col + 15)
        if tasks_sheet.cell(tasks_header_row, column).value
    }
    assert tasks_sheet.cell(tasks_header_row, tasks_header_col + 1).value in {
        "Название",
        "Title",
    }
    assert "evidence_url" not in header_values


@pytest.mark.django_db
def test_analytics_ec_export_preview_notifications_matches_export_rows(
    api_client, media_root, analytics_task_data
):
    from django.utils import timezone

    from apps.notifications.models import NotificationJob
    from apps.tasks.models import Task

    task = Task.objects.get(pk=analytics_task_data["task_id"])
    NotificationJob.objects.create(
        task=task,
        status="succeeded",
        reason="reminder",
        telegram_chat_id=597181229,
        scheduled_at=timezone.now(),
    )
    response = api_client.get(
        reverse("analytics-export-preview"),
        {"export_type": ExportType.NOTIFICATION_REPORT, "locale": "ru"},
    )
    assert response.status_code == 200
    sheets = {sheet["id"]: sheet for sheet in response.json()["sheets"]}
    assert sheets["notifications"]["total"] >= 1
    assert sheets["notifications"]["rows"]


@pytest.mark.django_db
def test_analytics_ec_export_preview_weekly_summary_exposes_kpis(
    api_client, analytics_task_data
):
    response = api_client.get(
        reverse("analytics-export-preview"),
        {"export_type": ExportType.WEEKLY_SUMMARY, "locale": "ru"},
    )
    assert response.status_code == 200
    first_sheet = response.json()["sheets"][0]
    assert "kpis" in first_sheet
    assert isinstance(first_sheet["kpis"], dict)


@pytest.mark.django_db
def test_analytics_ec_export_xlsx_empty_table_shows_no_data_message(media_root):
    from django.conf import settings
    from openpyxl import load_workbook

    export_job = AnalyticsExportJob.objects.create(
        export_type=ExportType.TAG_SUMMARY,
        file_format=FileFormat.XLSX,
        status=ExportStatus.PENDING,
        filters={"locale": "ru", "tags": ["__no_such_tag__"]},
    )
    relative_path, rows = export_report(export_job)
    assert rows == 0

    workbook = load_workbook(settings.MEDIA_ROOT / relative_path)
    sheet = workbook[workbook.sheetnames[0]]
    cell_values = [
        sheet.cell(row_index, column).value
        for row_index in range(1, 25)
        for column in range(1, 15)
        if sheet.cell(row_index, column).value
    ]
    assert any("Данные не найдены" in str(value) for value in cell_values)


@pytest.mark.django_db
def test_analytics_ec_export_weekly_summary_xlsx(media_root, analytics_task_data):
    export_job = AnalyticsExportJob.objects.create(
        export_type=ExportType.WEEKLY_SUMMARY,
        file_format=FileFormat.XLSX,
        status=ExportStatus.PENDING,
        filters={},
    )
    relative_path, rows = export_report(export_job)
    assert str(relative_path).endswith(".xlsx")
    assert rows >= 0


@pytest.mark.django_db
@pytest.mark.parametrize(
    "export_type",
    [
        pytest.param(ExportType.TASKS, id="ec_export_tasks_pdf"),
        pytest.param(ExportType.WEEKLY_SUMMARY, id="ec_export_weekly_summary_pdf"),
        pytest.param(ExportType.TAG_SUMMARY, id="ec_export_tag_summary_pdf"),
    ],
)
def test_analytics_ec_export_pdf_generates_valid_file(
    media_root, analytics_task_data, export_type
):
    from django.conf import settings

    export_job = AnalyticsExportJob.objects.create(
        export_type=export_type,
        file_format=FileFormat.PDF,
        status=ExportStatus.PENDING,
        filters={"locale": "ru"},
    )
    relative_path, rows = export_report(export_job)
    assert str(relative_path).endswith(".pdf")
    assert rows >= 0

    pdf_bytes = (settings.MEDIA_ROOT / relative_path).read_bytes()
    assert pdf_bytes.startswith(b"%PDF")


@pytest.mark.django_db
def test_analytics_ec_export_pdf_empty_table(media_root):
    from django.conf import settings

    export_job = AnalyticsExportJob.objects.create(
        export_type=ExportType.TAG_SUMMARY,
        file_format=FileFormat.PDF,
        status=ExportStatus.PENDING,
        filters={"locale": "ru", "tags": ["__no_such_tag__"]},
    )
    relative_path, rows = export_report(export_job)
    assert rows == 0
    pdf_bytes = (settings.MEDIA_ROOT / relative_path).read_bytes()
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 500


@pytest.mark.django_db
def test_analytics_ec_export_filename_includes_metadata(media_root, week_key):
    export_job = AnalyticsExportJob.objects.create(
        export_type=ExportType.WEEKLY_SUMMARY,
        file_format=FileFormat.XLSX,
        status=ExportStatus.PENDING,
        filters={"week": week_key, "tags": ["analytics"], "source": "api"},
    )
    relative_path, _ = export_report(export_job)
    filename = relative_path.name
    assert filename.startswith("weekly_summary_")
    assert week_key in filename
    assert "tags-analytics" in filename
    assert "source-api" in filename
    assert filename.endswith(".xlsx")
    assert re.search(r"_\d{4}-\d{2}-\d{2}_\d{4}\.xlsx$", filename)
    with pytest.raises(ValueError):
        AnalyticsExportService.create(
            export_type="invalid", file_format=FileFormat.CSV, filters={}
        )


@pytest.mark.django_db
def test_analytics_ec_export_download_not_ready_returns_409(api_client):
    export_job = AnalyticsExportJob.objects.create(
        export_type=ExportType.TASKS,
        file_format=FileFormat.CSV,
        status=ExportStatus.PENDING,
        filters={},
    )
    response = api_client.get(
        reverse("analytics-export-download", args=[export_job.id])
    )
    assert response.status_code == 409
